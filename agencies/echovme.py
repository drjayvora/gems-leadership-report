"""EchoVME spend sheet -> spend/echovme.csv (the standard spend format, see spend/TEMPLATE.csv).

EchoVME sends one Google Sheet with a tab per month and side-by-side blocks per platform, campus and
objective. BLOCKS says where each block's columns are (0-based). If they change the layout, edit BLOCKS.
Usage: python3 agencies/echovme.py [path-to-xlsx]   (default: download the sheet whose ID is in env ECHOVME_SHEET)
"""
import csv, datetime, io, json, os, sys, urllib.request
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
CFG = json.load(open(os.path.join(ROOT, "config.json")))
COIMBATORE, RAIPUR = "GEMS Modern Academy - Coimbatore", "GEMS Millennium School - Raipur"
VIKHROLI, KALYAN = "GEMS Millennium School - Vikhroli", "GEMS Millennium School - Kalyan"

# tab: [(school, channel, objective, date_col, spend_col, impressions_col, reach_col, leads_col)]
BLOCKS = {
    "Sept": [(RAIPUR, "Meta", "Branding", 0, 1, 3, 2, None), (COIMBATORE, "Meta", "Branding", 0, 6, 8, 7, None),
             (RAIPUR, "Google", "Branding", 11, 12, 13, None, None), (COIMBATORE, "Google", "Branding", 11, 16, 17, None, None)],
    # Oct layout (changed ~20 Oct): Meta blocks Raipur, Coimbatore branding, Coimbatore lead gen, Vikhroli, Kalyan;
    # Google Coimbatore branding + lead gen from column 25.
    "Oct": [(RAIPUR, "Meta", "Branding", 0, 1, 3, 2, None), (COIMBATORE, "Meta", "Branding", 0, 6, 8, 7, None),
            (COIMBATORE, "Meta", "Lead gen", 0, 11, None, None, 12),
            (VIKHROLI, "Meta", "Branding", 0, 15, 17, 16, None), (KALYAN, "Meta", "Branding", 0, 20, 22, 21, None),
            (COIMBATORE, "Google", "Branding", 25, 26, 27, None, None), (COIMBATORE, "Google", "Lead gen", 25, 29, None, None, 30)],
}


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("₹", "").replace(",", "").strip())
    except ValueError:
        return None


def rows(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    for tab, blocks in BLOCKS.items():
        if tab not in wb.sheetnames:
            print("missing tab", tab); continue
        for r in wb[tab].iter_rows(min_row=5, values_only=True):
            for school, ch, obj, dc, sc, ic, rc, lc in blocks:
                d, spend = r[dc] if dc < len(r) else None, num(r[sc]) if sc < len(r) else None
                if not isinstance(d, datetime.datetime) or not spend:
                    continue
                g = lambda c: (num(r[c]) or 0) if c is not None and c < len(r) else ""
                yield dict(date=d.date().isoformat(), agency="EchoVME", school=school, channel=ch, objective=obj,
                           spend=round(spend, 2), impressions=g(ic), reach=g(rc), agency_leads=g(lc))


if __name__ == "__main__":
    url = f"https://docs.google.com/spreadsheets/d/{os.environ.get(CFG['agencies']['EchoVME']['sheet_env'], '')}/export?format=xlsx"
    src = sys.argv[1] if len(sys.argv) > 1 else io.BytesIO(urllib.request.urlopen(url, timeout=60).read())
    out = list(rows(src))
    with open(os.path.join(ROOT, "spend", "echovme.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["date", "agency", "school", "channel", "objective", "spend", "impressions", "reach", "agency_leads"])
        w.writeheader(); w.writerows(sorted(out, key=lambda x: (x["date"], x["school"], x["channel"], x["objective"])))
    print(len(out), "rows, spend", round(sum(x["spend"] for x in out)))
